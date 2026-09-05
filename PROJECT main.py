import datetime
import os
from fastapi import FastAPI, APIRouter, Depends, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import create_engine, Column, Integer, String, DateTime, Text, Boolean
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker, Session
from pydantic import BaseModel
from dotenv import load_dotenv
import pandas as pd
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

# --------------------- Load Environment ---------------------
load_dotenv()
SQLALCHEMY_DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./app.db")

# --------------------- Database Setup ---------------------
engine = create_engine(SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# --------------------- SQLAlchemy Models ---------------------
class LessonLearned(Base):
    __tablename__ = "lessons_learned"
    id = Column(Integer, primary_key=True)
    project_id = Column(Integer)
    title = Column(String)
    lesson = Column(Text)
    recommendation = Column(Text)
    validated = Column(Boolean, default=False)
    created_at = Column(DateTime, default=datetime.datetime.utcnow)

class Evidence(Base):
    __tablename__ = "evidence"
    id = Column(Integer, primary_key=True)
    indicator = Column(String)
    source = Column(String)
    file_path = Column(String)
    verified = Column(Boolean, default=False)
    created_at = Column(DateTime, default=datetime.datetime.utcnow)

class DecisionLog(Base):
    __tablename__ = "decision_logs"
    id = Column(Integer, primary_key=True)
    decision = Column(Text)
    rationale = Column(Text)
    owner = Column(String)
    status = Column(String)
    created_at = Column(DateTime, default=datetime.datetime.utcnow)

Base.metadata.create_all(bind=engine)

# --------------------- Pydantic Schemas ---------------------
class LessonCreate(BaseModel):
    project_id: int
    title: str
    lesson: str
    recommendation: str

class EvidenceCreate(BaseModel):
    indicator: str
    source: str
    file_path: str

class DecisionCreate(BaseModel):
    decision: str
    rationale: str
    owner: str

# --------------------- RBM Engines ---------------------
class TheoryOfChangeEngine:
    def generate(self, project):
        sector = project.get("sector", "general").lower()
        templates = {
            "education": {
                "outcomes": [
                    {"name": "Increased student enrollment"},
                    {"name": "Improved learning outcomes"}
                ],
                "outputs": [
                    {"name": "Teacher training delivered"},
                    {"name": "Learning materials distributed"}
                ],
                "activities": ["Train teachers", "Distribute books", "Monitor attendance"]
            },
            "health": {
                "outcomes": [
                    {"name": "Reduced disease incidence"},
                    {"name": "Improved maternal health"}
                ],
                "outputs": [
                    {"name": "Health clinics equipped"},
                    {"name": "Community health workers trained"}
                ],
                "activities": ["Supply medicines", "Conduct awareness campaigns", "Train health staff"]
            }
        }
        template = templates.get(sector, {
            "outcomes": [{"name": "Improved capacity"}, {"name": "Improved access"}],
            "outputs": [{"name": "Training delivered"}, {"name": "Services strengthened"}],
            "activities": ["Training", "Coaching", "Advocacy"]
        })
        return {
            "problem": project.get("problem", ""),
            "impact": project.get("impact", ""),
            "outcomes": template["outcomes"],
            "outputs": template["outputs"],
            "activities": template["activities"]
        }

class ResultsFrameworkBuilder:
    def build(self, toc):
        return {
            "impact": toc.get("impact", ""),
            "outcomes": toc.get("outcomes", []),
            "outputs": toc.get("outputs", [])
        }

class LogframeBuilder:
    def generate(self, framework):
        logframe = []
        logframe.append({
            "level": "Impact",
            "statement": framework.get("impact", ""),
            "indicator": "Impact indicator",
            "baseline": 0,
            "target": 100
        })
        for outcome in framework.get("outcomes", []):
            logframe.append({
                "level": "Outcome",
                "statement": outcome.get("name", ""),
                "indicator": f"% improvement in {outcome.get('name', '')}",
                "baseline": 0,
                "target": 80
            })
        return logframe

class IndicatorEngine:
    def generate_outcome_indicators(self, outcomes):
        indicators = []
        for outcome in outcomes:
            indicators.append({
                "result": outcome.get("name", ""),
                "indicator": f"Percentage of beneficiaries reporting improvements in {outcome.get('name', '')}",
                "unit": "%",
                "baseline": 0,
                "target": 80
            })
        return indicators

    def generate_output_indicators(self, outputs):
        indicators = []
        for output in outputs:
            indicators.append({
                "result": output.get("name", ""),
                "indicator": f"Number of {output.get('name', '')} completed",
                "unit": "Count",
                "baseline": 0,
                "target": 100
            })
        return indicators

class RiskEngine:
    def generate(self, sector="general"):
        risks_map = {
            "education": [
                {"risk": "Teacher absenteeism", "likelihood": "Medium", "impact": "High", "mitigation": "Regular monitoring"},
                {"risk": "Low parental engagement", "likelihood": "Medium", "impact": "Medium", "mitigation": "Community outreach"}
            ],
            "health": [
                {"risk": "Disease outbreak", "likelihood": "Medium", "impact": "High", "mitigation": "Emergency preparedness"},
                {"risk": "Supply chain disruption", "likelihood": "High", "impact": "High", "mitigation": "Diversify suppliers"}
            ]
        }
        return risks_map.get(sector, [
            {"risk": "Political instability", "likelihood": "Medium", "impact": "High", "mitigation": "Stakeholder engagement"},
            {"risk": "Funding delay", "likelihood": "Medium", "impact": "Medium", "mitigation": "Contingency reserves"}
        ])

class StakeholderEngine:
    def map_stakeholders(self, sector="general"):
        stakeholders_map = {
            "education": [
                {"stakeholder": "Ministry of Education", "interest": "High", "influence": "High"},
                {"stakeholder": "School principals", "interest": "High", "influence": "Medium"},
                {"stakeholder": "Parents", "interest": "High", "influence": "Low"}
            ],
            "health": [
                {"stakeholder": "Ministry of Health", "interest": "High", "influence": "High"},
                {"stakeholder": "Hospital administrators", "interest": "High", "influence": "Medium"},
                {"stakeholder": "Community leaders", "interest": "Medium", "influence": "Medium"}
            ]
        }
        return stakeholders_map.get(sector, [
            {"stakeholder": "Government", "interest": "High", "influence": "High"},
            {"stakeholder": "Community", "interest": "High", "influence": "Medium"}
        ])

class AssumptionsEngine:
    def build(self, sector="general"):
        assumptions_map = {
            "education": [
                "Teachers have adequate qualifications",
                "Students attend regularly",
                "School facilities remain available"
            ],
            "health": [
                "Community accepts health interventions",
                "Supply chain functions",
                "Government prioritises health"
            ]
        }
        return assumptions_map.get(sector, [
            "Target group participates",
            "Partners remain engaged",
            "Resources remain available",
            "Enabling environment remains stable"
        ])

# --------------------- Knowledge Engines ---------------------
class LearningEngine:
    def classify_lesson(self, lesson_text):
        return {"category": "Implementation", "theme": "Community Engagement", "status": "Emerging Observation"}

    def convert_to_learning_brief(self, lesson):
        return {
            "title": lesson.get("title", ""),
            "key_learning": lesson.get("lesson", ""),
            "recommended_action": lesson.get("recommendation", "")
        }

class EvidenceEngine:
    def validate_evidence(self, source):
        return {"source": source, "verification_status": "Pending", "review_required": True}

class KnowledgeSearch:
    def search(self, keyword):
        # In a real implementation, query the database; here a placeholder
        return {"keyword": keyword, "results": []}

class CaseStudyBuilder:
    def generate(self, project):
        return {
            "title": project.get("title", ""),
            "context": project.get("problem", ""),
            "intervention": project.get("solution", ""),
            "results": project.get("results", ""),
            "lesson": project.get("lesson", "")
        }

class DecisionLogEngine:
    def record(self, decision, rationale):
        return {"decision": decision, "rationale": rationale, "status": "Approved"}

# --------------------- Reporting Engines ---------------------
class IndicatorTracker:
    def calculate_progress(self, baseline, current, target):
        if target == baseline:
            return 0
        return round(((current - baseline) / (target - baseline)) * 100, 2)

class DashboardEngine:
    def build_dashboard(self, indicators):
        dashboard = []
        for i in indicators:
            progress = i.get("progress", 0)
            dashboard.append({
                "indicator": i.get("indicator", ""),
                "progress": progress,
                "status": self.status(progress)
            })
        return dashboard

    def status(self, value):
        if value >= 80:
            return "Green"
        if value >= 50:
            return "Amber"
        return "Red"

class ReportEngine:
    def generate_quarterly_report(self, project):
        return {
            "executive_summary": "Quarterly progress summary",
            "outputs_delivered": project.get("outputs", []),
            "outcomes_progress": project.get("outcomes", []),
            "risks": project.get("risks", []),
            "lessons": project.get("lessons", []),
            "next_steps": project.get("next_steps", [])
        }

class LearningReportEngine:
    def generate(self, lessons):
        report = []
        for lesson in lessons:
            report.append({
                "title": lesson.get("title", ""),
                "lesson": lesson.get("lesson", ""),
                "action": lesson.get("recommendation", "")
            })
        return report

class ExcelExporter:
    def export_logframe(self, data, filename):
        df = pd.DataFrame(data)
        df.to_excel(filename, index=False)
        return filename

class WordExporter:
    def export_report(self, report, filename):
        doc = Document()
        doc.add_heading('Quarterly Report', 0)
        for key, value in report.items():
            doc.add_heading(key, level=1)
            doc.add_paragraph(str(value))
        doc.save(filename)
        return filename

class PDFExporter:
    def export(self, report, filename):
        doc = SimpleDocTemplate(filename)
        styles = getSampleStyleSheet()
        elements = []
        elements.append(Paragraph("Donor Report", styles["Title"]))
        for k, v in report.items():
            elements.append(Paragraph(f"<b>{k}</b>", styles["Heading2"]))
            elements.append(Paragraph(str(v), styles["BodyText"]))
        doc.build(elements)
        return filename

# --------------------- FastAPI App ---------------------
app = FastAPI(title="RBM & Knowledge System", version="1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --------------------- RBM Router ---------------------
rbm_router = APIRouter(prefix="/rbm", tags=["RBM Engine"])

@rbm_router.post("/generate")
def generate_rbm(project: dict):
    sector = project.get("sector", "general")
    toc_engine = TheoryOfChangeEngine()
    framework_builder = ResultsFrameworkBuilder()
    logframe_builder = LogframeBuilder()
    indicator_engine = IndicatorEngine()
    risk_engine = RiskEngine()
    assumptions_engine = AssumptionsEngine()
    stakeholder_engine = StakeholderEngine()

    toc = toc_engine.generate(project)
    framework = framework_builder.build(toc)
    logframe = logframe_builder.generate(framework)
    outcome_indicators = indicator_engine.generate_outcome_indicators(framework["outcomes"])
    output_indicators = indicator_engine.generate_output_indicators(framework["outputs"])

    return {
        "theory_of_change": toc,
        "results_framework": framework,
        "logframe": logframe,
        "outcome_indicators": outcome_indicators,
        "output_indicators": output_indicators,
        "risks": risk_engine.generate(sector),
        "assumptions": assumptions_engine.build(sector),
        "stakeholders": stakeholder_engine.map_stakeholders(sector)
    }

# --------------------- Knowledge Router ---------------------
knowledge_router = APIRouter(prefix="/knowledge", tags=["Knowledge Management"])

@knowledge_router.post("/lessons")
def create_lesson(lesson: LessonCreate, db: Session = Depends(get_db)):
    record = LessonLearned(**lesson.dict())
    db.add(record)
    db.commit()
    db.refresh(record)
    return record

@knowledge_router.get("/lessons")
def list_lessons(db: Session = Depends(get_db)):
    return db.query(LessonLearned).all()

@knowledge_router.post("/evidence")
def create_evidence(item: EvidenceCreate, db: Session = Depends(get_db)):
    record = Evidence(**item.dict())
    db.add(record)
    db.commit()
    db.refresh(record)
    return record

@knowledge_router.post("/decisions")
def create_decision(item: DecisionCreate, db: Session = Depends(get_db)):
    record = DecisionLog(**item.dict())
    db.add(record)
    db.commit()
    db.refresh(record)
    return record

# Additional knowledge endpoints (using engines)
@knowledge_router.post("/classify-lesson")
def classify_lesson(lesson_text: str):
    engine = LearningEngine()
    return engine.classify_lesson(lesson_text)

@knowledge_router.post("/learning-brief")
def learning_brief(lesson: dict):
    engine = LearningEngine()
    return engine.convert_to_learning_brief(lesson)

@knowledge_router.post("/validate-evidence")
def validate_evidence(source: str):
    engine = EvidenceEngine()
    return engine.validate_evidence(source)

@knowledge_router.get("/search")
def search_knowledge(keyword: str):
    engine = KnowledgeSearch()
    return engine.search(keyword)

@knowledge_router.post("/case-study")
def generate_case_study(project: dict):
    engine = CaseStudyBuilder()
    return engine.generate(project)

# --------------------- Reporting Router ---------------------
reporting_router = APIRouter(prefix="/reporting", tags=["Reporting"])

@reporting_router.post("/quarterly")
def quarterly_report(project: dict):
    engine = ReportEngine()
    return engine.generate_quarterly_report(project)

@reporting_router.post("/indicator-progress")
def indicator_progress(baseline: int, current: int, target: int):
    tracker = IndicatorTracker()
    return {"progress": tracker.calculate_progress(baseline, current, target)}

@reporting_router.post("/dashboard")
def build_dashboard(indicators: list):
    engine = DashboardEngine()
    return engine.build_dashboard(indicators)

@reporting_router.post("/learning-report")
def learning_report(lessons: list):
    engine = LearningReportEngine()
    return engine.generate(lessons)

@reporting_router.post("/export-excel")
def export_excel(data: list, filename: str = "logframe.xlsx"):
    exporter = ExcelExporter()
    return {"file": exporter.export_logframe(data, filename)}

@reporting_router.post("/export-word")
def export_word(report: dict, filename: str = "report.docx"):
    exporter = WordExporter()
    return {"file": exporter.export_report(report, filename)}

@reporting_router.post("/export-pdf")
def export_pdf(report: dict, filename: str = "report.pdf"):
    exporter = PDFExporter()
    return {"file": exporter.export(report, filename)}

# --------------------- Include Routers ---------------------
app.include_router(rbm_router)
app.include_router(knowledge_router)
app.include_router(reporting_router)

# --------------------- Root Endpoint ---------------------
@app.get("/")
def root():
    return {"message": "RBM System API"}

# --------------------- Run ---------------------
if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)